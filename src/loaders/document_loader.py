from pathlib import Path

from langchain_core.documents import Document
from langchain_community.document_loaders import (
    PyPDFLoader,
    CSVLoader,
    UnstructuredWordDocumentLoader,
)


def load_txt(file_path: str) -> list[Document]:
    """Load a TXT file."""
    text = Path(file_path).read_text(encoding="utf-8")

    return [
        Document(
            page_content=text,
            metadata={
                "source": str(file_path),
                "file_type": "txt",
            },
        )
    ]


def load_pdf(file_path: str) -> list[Document]:
    """Load a PDF file."""
    loader = PyPDFLoader(file_path)
    documents = loader.load()

    for document in documents:
        document.metadata["file_type"] = "pdf"

    return documents


def load_csv(file_path: str) -> list[Document]:
    """Load a CSV file."""
    loader = CSVLoader(file_path)
    documents = loader.load()

    for document in documents:
        document.metadata["file_type"] = "csv"

    return documents


def load_docx(file_path: str) -> list[Document]:
    """Load a DOCX file."""
    loader = UnstructuredWordDocumentLoader(file_path)
    documents = loader.load()

    for document in documents:
        document.metadata["file_type"] = "docx"

    return documents


def load_xlsx(file_path: str) -> list[Document]:
    """Load an Excel XLSX file."""
    from openpyxl import load_workbook

    workbook = load_workbook(file_path, data_only=True)

    documents = []

    for sheet in workbook.worksheets:
        rows = []

        for row in sheet.iter_rows(values_only=True):
            values = [str(value) for value in row if value is not None]

            if values:
                rows.append(" | ".join(values))

        if rows:
            text = "\n".join(rows)

            documents.append(
                Document(
                    page_content=text,
                    metadata={
                        "source": str(file_path),
                        "file_type": "xlsx",
                        "sheet": sheet.title,
                    },
                )
            )

    return documents


def load_document(file_path: str) -> list[Document]:
    """Load a document based on its file extension."""

    extension = Path(file_path).suffix.lower()

    if extension == ".txt":
        return load_txt(file_path)

    if extension == ".pdf":
        return load_pdf(file_path)

    if extension == ".csv":
        return load_csv(file_path)

    if extension == ".docx":
        return load_docx(file_path)

    if extension == ".xlsx":
        return load_xlsx(file_path)

    raise ValueError(
        f"Unsupported file type: {extension}. "
        "Supported types: .txt, .pdf, .csv, .xlsx, .docx"
    )


def load_documents(folder_path: str) -> list[Document]:
    """Load all supported documents from a folder."""

    folder = Path(folder_path)

    if not folder.exists():
        raise FileNotFoundError(
            f"Document folder does not exist: {folder_path}"
        )

    documents = []

    supported_extensions = {
        ".txt",
        ".pdf",
        ".csv",
        ".xlsx",
        ".docx",
    }

    for file_path in folder.iterdir():

        if not file_path.is_file():
            continue

        if file_path.suffix.lower() not in supported_extensions:
            continue

        print(f"Loading: {file_path.name}")

        file_documents = load_document(str(file_path))

        documents.extend(file_documents)

    return documents